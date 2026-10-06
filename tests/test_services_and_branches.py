import re
import sqlite3
from zipfile import is_zipfile

from controllers.biblioteca_controller import BibliotecaController
from database.db_manager import DatabaseManager
from openpyxl import load_workbook
from services.backup_service import BackupService
from services.excel_service import ExcelService
from services.pdf_service import PDFService


def test_backup_excel_y_pdf_se_generan_y_rotan(tmp_path):
    database = DatabaseManager(tmp_path / "central.db")
    database.initialize()
    backups = BackupService(database.database_path, tmp_path / "backups", max_backups=7)
    for _ in range(8):
        exito, resultado = backups.crear_backup()
        assert exito is True, resultado
    copias = list((tmp_path / "backups").glob("biblioteca_*.db"))
    assert len(copias) == 7
    with sqlite3.connect(resultado) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"

    ExcelService().exportar({"Inventario": [{"titulo": "Libro", "estado": "recibido"}]}, tmp_path / "inventario.xlsx")
    assert is_zipfile(tmp_path / "inventario.xlsx")
    respaldo_excel = tmp_path / "respaldo_completo.xlsx"
    ExcelService().exportar_base_datos(database.engine, respaldo_excel, administrador=True)
    libro_excel = load_workbook(respaldo_excel, read_only=True)
    assert {"libros", "usuarios", "movimientos", "bulto_libros"} <= set(libro_excel.sheetnames)
    columnas_usuario = next(libro_excel["usuarios"].iter_rows(values_only=True))
    assert "username" in columnas_usuario
    assert "password_hash" not in columnas_usuario
    assert "salt" not in columnas_usuario
    libro_excel.close()
    formula_excel = tmp_path / "formula_segura.xlsx"
    valores_peligrosos = ["=1+1", " +1+1", "-1+1", "@SUM(A1:A2)", "\t=1+1", "\r=1+1"]
    ExcelService().exportar({
        "Prueba": [{"texto": valor, "numero": 12.5} for valor in valores_peligrosos],
    }, formula_excel)
    libro_seguro = load_workbook(formula_excel, data_only=False)
    for fila, valor in enumerate(valores_peligrosos, start=2):
        celda = libro_seguro["Prueba"].cell(fila, 1)
        assert celda.data_type == "s"
        assert celda.value == f"'{valor}".replace("\r", "\n")
        assert libro_seguro["Prueba"].cell(fila, 2).value == 12.5
    libro_seguro.close()
    pdf = PDFService()
    bulto = {"codigo_envio": "ENV-20260930-001", "destino": "Sucursal", "fecha": "2026-09-30", "cantidad_libros": 1}
    libros = [
        {"titulo": f"Libro {indice} <de prueba> & otros", "autor": "Autora & co.", "cota": "863.6\nA123\n2024", "numero_registro": f"REG-2026-{indice:05d}"}
        for indice in range(1, 5)
    ]
    for ruta in (tmp_path / "cotas.pdf", tmp_path / "fichas.pdf", tmp_path / "envio.pdf", tmp_path / "mini.pdf"):
        if ruta.name == "cotas.pdf":
            pdf.generar_cotas(libros, ruta, ancho_cm=1.6, alto_cm=4.2)
        elif ruta.name == "fichas.pdf":
            pdf.generar_fichas(libros, ruta)
        elif ruta.name == "envio.pdf":
            pdf.generar_documento_envio(bulto, libros, ruta)
        else:
            pdf.generar_mini_ficha(bulto, ruta)
        assert ruta.read_bytes().startswith(b"%PDF-")
    fichas_unica_pagina = (tmp_path / "fichas.pdf").read_bytes()
    assert len(re.findall(rb"/Type\s*/Page\b", fichas_unica_pagina)) == 1
    control = tmp_path / "control_envio.pdf"
    pdf.generar_control_envio_snbp(
        bulto,
        [
            {
                **libros[0],
                "organismo": "Librería Central",
                "procedencia": "compra",
                "numero_volumenes": 2,
                "precio_unitario": "125.50",
            },
            {**libros[1], "procedencia": "donacion"},
        ],
        control,
    )
    assert control.read_bytes().startswith(b"%PDF-")
    nota = tmp_path / "nota_entrega.pdf"
    pdf.generar_nota_entrega(
        {
            **bulto,
            "destino": "Biblioteca Pública Sur",
            "direccion": "Avenida Bolívar",
            "municipio": "Angostura del Orinoco",
            "encargada": "María Pérez",
            "cantidad_volumenes": 42,
        },
        {"nombre": "María Pérez", "cedula": "V-12345678"},
        nota,
    )
    assert nota.read_bytes().startswith(b"%PDF-")
    matriz = tmp_path / "matriz_sucursales.pdf"
    pdf.generar_matriz_control_sucursales(
        {
            "titulo": "Historia de Venezuela",
            "autor": "Rómulo Gallegos",
            "numero_registro": "REG-2026-00001",
            "cota": "863.6",
            "estado": "distribuido",
            "filas": [
                {
                    "biblioteca": f"Biblioteca {indice}",
                    "ingreso": indice == 0,
                    "disponibilidad": indice == 0,
                    "prestamo": False,
                }
                for indice in range(26)
            ],
        },
        matriz,
    )
    assert matriz.read_bytes().startswith(b"%PDF-")
    resumen_areas = tmp_path / "resumen_areas.pdf"
    pdf.generar_resumen_distribucion_bibliotecas(
        {
            "anio": 2026,
            "sucursales": [{
                "biblioteca": "Biblioteca Central 'Rómulo Gallegos'",
                "categorias": {
                    **{f"{area:03d}": {"T": 0, "V": 0} for area in range(0, 1000, 100)},
                    "Biografías": {"T": 0, "V": 0},
                    "Pub. Oficiales": {"T": 0, "V": 0},
                    "Pub. Periódicas": {"T": 0, "V": 0},
                    "No Bibliográfico": {"T": 0, "V": 0},
                    "Total": {"T": 2, "V": 4},
                    "Total General": {"T": 2, "V": 4},
                    "800": {"T": 2, "V": 4},
                },
            }],
            "totales_generales": {
                **{f"{area:03d}": {"T": 0, "V": 0} for area in range(0, 1000, 100)},
                "Biografías": {"T": 0, "V": 0},
                "Pub. Oficiales": {"T": 0, "V": 0},
                "Pub. Periódicas": {"T": 0, "V": 0},
                "No Bibliográfico": {"T": 0, "V": 0},
                "Total": {"T": 2, "V": 4},
                "Total General": {"T": 2, "V": 4},
            },
            "nota_clasificacion": "Las publicaciones oficiales no están diferenciadas en la base.",
        },
        resumen_areas,
    )
    assert resumen_areas.read_bytes().startswith(b"%PDF-")
    assert b"792 612" in resumen_areas.read_bytes()
    import pytest
    with pytest.raises(ValueError, match="nombre y la cédula"):
        pdf.generar_nota_entrega(
            {**bulto, "cantidad_volumenes": 1}, {"nombre": "", "cedula": ""},
            tmp_path / "nota_invalida.pdf",
        )
    ficha_individual = tmp_path / "ficha_individual.pdf"
    pdf.generar_fichas([{
        **libros[0],
        "observaciones": "Ejemplar para consulta",
        "numero_volumenes": 2,
    }], ficha_individual)
    contenido_ficha_individual = ficha_individual.read_bytes()
    assert contenido_ficha_individual.startswith(b"%PDF-")
    assert len(re.findall(rb"/Type\s*/Page\b", contenido_ficha_individual)) == 1
    fichas_siguiente_hoja = tmp_path / "cinco_fichas.pdf"
    pdf.generar_fichas([*libros, libros[0]], fichas_siguiente_hoja)
    contenido_cinco_fichas = fichas_siguiente_hoja.read_bytes()
    assert contenido_cinco_fichas.startswith(b"%PDF-")
    assert len(re.findall(rb"/Type\s*/Page\b", contenido_cinco_fichas)) == 2
    with pytest.raises(ValueError, match="al menos un libro"):
        pdf.generar_fichas([], tmp_path / "sin_fichas.pdf")
    reporte = tmp_path / "reporte.pdf"
    pdf.generar_reporte("Reporte & catálogo", [("Libros", ["Título", "Registro"], [["Libro <A>", "REG-1"]])], reporte)
    assert reporte.read_bytes().startswith(b"%PDF-")
    database.close()


def test_sucursales_pueden_crearse_y_desactivarse_sin_desactivar_central(tmp_path):
    database = DatabaseManager(tmp_path / "central.db")
    database.initialize()
    controller = BibliotecaController(database)
    creado, mensaje = controller.crear("Sucursal Sur", "Avenida 1", "Encargada", "555-0100", "sur@example.org")
    assert creado is True, mensaje
    sucursal = next(item for item in controller.listar() if item.nombre == "Sucursal Sur")
    assert controller.desactivar(sucursal.id)[0] is True
    assert all(item.nombre != "Sucursal Sur" for item in controller.listar())
    central = next(item for item in controller.listar() if item.nombre == "Biblioteca Central Rómulo Gallegos")
    assert controller.desactivar(central.id)[0] is False
    database.close()