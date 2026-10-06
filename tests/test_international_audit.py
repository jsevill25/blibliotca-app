import sqlite3

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError

from controllers.auth_controller import AuthController
from controllers.backup_controller import BackupController
from database.db_manager import DatabaseManager
from database.models import AuditoriaLog, User
from services.excel_service import ExcelService
from services.pdf_service import PDFService
from web_preview import PreviewHandler


def test_pending_password_blocks_operational_auth_controller_actions(tmp_path):
    database = DatabaseManager(tmp_path / "auth.sqlite")
    database.initialize()
    auth = AuthController(database)
    assert auth.autenticar("admin", "clave-incorrecta")[0] is False
    valido, admin = auth.autenticar("admin", "admin123")

    assert valido is True
    assert admin.debe_cambiar_clave is True
    assert auth.acceso_operativo_permitido(admin.id) is False
    creado, _ = auth.crear_usuario(
        "operador", "Operador de prueba", "ClaveOperador2026", "bibliotecario", admin.id,
    )
    assert creado is False

    assert auth.cambiar_clave(admin.id, "admin123", "ClaveAdministrativa2026")[0] is True
    assert auth.acceso_operativo_permitido(admin.id) is True
    assert auth.crear_usuario(
        "operador", "Operador de prueba", "ClaveOperador2026", "bibliotecario", admin.id,
    )[0] is True

    with database.session() as session:
        acciones = list(session.scalars(select(AuditoriaLog.accion)))
    assert acciones.count("LOGIN_EXITOSO") == 1
    assert acciones.count("LOGIN_FALLIDO") == 1
    assert acciones.count("CAMBIO_CLAVE") == 1
    assert acciones.count("ACCION_ADMIN") == 1
    database.close()


def test_excel_neutralizes_formula_prefixes_and_omits_secret_columns(tmp_path):
    database = DatabaseManager(tmp_path / "export.sqlite")
    database.initialize()
    secretos = ("hash-que-no-debe-exportarse", "salt-que-no-debe-exportarse")
    with database.session() as session:
        usuario = session.scalar(select(User).where(User.username == "admin"))
        secretos_cuenta = (usuario.password_hash, usuario.salt)
    ruta = tmp_path / "seguro.xlsx"
    ExcelService().exportar({
        "Reporte": [
            {"texto": valor, "password_hash": secretos[0], "salt": secretos[1]}
            for valor in ("=1+1", " +1+1", "-1+1", "@SUM(A1:A2)", "\t=1+1")
        ],
    }, ruta)

    workbook = load_workbook(ruta, data_only=False)
    sheet = workbook["Reporte"]
    assert "password_hash" not in [cell.value for cell in sheet[1]]
    assert "salt" not in [cell.value for cell in sheet[1]]
    for row in range(2, 7):
        assert sheet.cell(row, 1).data_type == "s"
        assert sheet.cell(row, 1).value.startswith("'")
    workbook.close()

    ruta_integral = tmp_path / "integral.xlsx"
    with pytest.raises(PermissionError, match="autorización administrativa"):
        ExcelService().exportar_base_datos(database.engine, ruta_integral)
    ExcelService().exportar_base_datos(database.engine, ruta_integral, administrador=True)
    workbook = load_workbook(ruta_integral, read_only=True)
    users = workbook["usuarios"]
    headers = next(users.iter_rows(values_only=True))
    assert "password_hash" not in headers
    assert "salt" not in headers
    assert not any(
        secret in str(value)
        for row in users.iter_rows(values_only=True)
        for value in row
        for secret in secretos_cuenta
    )
    workbook.close()
    database.close()


def test_pdf_reports_drop_sensitive_fields_before_rendering(tmp_path, monkeypatch):
    import services.pdf_service as pdf_module

    contenido_tablas = []
    tabla_original = pdf_module.Table

    def capturar_tabla(data, *args, **kwargs):
        contenido_tablas.extend(data)
        return tabla_original(data, *args, **kwargs)

    monkeypatch.setattr(pdf_module, "Table", capturar_tabla)
    destino = tmp_path / "reporte.pdf"
    PDFService().generar_reporte(
        "Reporte general",
        [(
            "Usuarios",
            ["username", "password_hash", "salt"],
            [["bibliotecario", "hash-privado", "salt-privada"]],
        )],
        destino,
    )

    assert destino.read_bytes().startswith(b"%PDF-")
    contenido = repr(contenido_tablas)
    assert "password_hash" not in contenido
    assert "salt" not in contenido
    assert "hash-privado" not in contenido
    assert "salt-privada" not in contenido


def test_sqlite_foreign_keys_are_active_and_reject_invalid_audit_user(tmp_path):
    database = DatabaseManager(tmp_path / "foreign-keys.sqlite")
    database.initialize()
    with database.engine.connect() as connection:
        assert connection.scalar(text("PRAGMA foreign_keys")) == 1

    with pytest.raises(IntegrityError):
        with database.session() as session:
            session.add(AuditoriaLog(
                usuario_id=987654, accion="LOGIN_EXITOSO",
                origen="prueba", detalle="Debe fallar por FK.",
            ))
            session.flush()
    database.close()


def test_backup_verifies_integrity_and_records_audit_event(tmp_path):
    database = DatabaseManager(tmp_path / "backup.sqlite")
    database.initialize()
    respaldo = BackupController(database, tmp_path / "backups")
    correcto, ruta = respaldo.crear_backup(usuario_id=1, origen="pytest")

    assert correcto is True
    with sqlite3.connect(ruta) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    with database.session() as session:
        evento = session.scalar(select(AuditoriaLog).where(AuditoriaLog.accion == "RESPALDO_DB"))
        assert evento is not None
        assert evento.usuario_id == 1
        assert evento.origen == "pytest"
    database.close()


def test_session_cookie_has_strict_attributes_and_secure_in_production(monkeypatch):
    monkeypatch.delenv("BLIBLIOTECA_ENV", raising=False)
    monkeypatch.delenv("APP_ENV", raising=False)
    monkeypatch.delenv("BLIBLIOTECA_COOKIE_SECURE", raising=False)
    cookie = PreviewHandler._session_cookie("token")
    assert "HttpOnly" in cookie
    assert "SameSite=Strict" in cookie
    assert "Secure" not in cookie

    monkeypatch.setenv("BLIBLIOTECA_ENV", "production")
    cookie_production = PreviewHandler._session_cookie("", expired=True)
    assert "HttpOnly" in cookie_production
    assert "SameSite=Strict" in cookie_production
    assert "Secure" in cookie_production
    assert "Max-Age=0" in cookie_production
