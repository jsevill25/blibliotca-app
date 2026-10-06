from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import MetaData


class ExcelService:
    SENSITIVE_COLUMNS = {"password_hash", "salt"}

    @classmethod
    def _columna_exportable(cls, nombre: str) -> bool:
        return nombre.strip().casefold() not in cls.SENSITIVE_COLUMNS

    def exportar(self, hojas: dict[str, list[dict]], destino: str | Path) -> Path:
        return self._escribir(hojas, destino)

    def exportar_base_datos(
        self, engine, destino: str | Path, *, administrador: bool = False,
    ) -> Path:
        if not administrador:
            raise PermissionError("La exportación integral requiere autorización administrativa.")
        metadata = MetaData()
        metadata.reflect(bind=engine)
        hojas = {}
        with engine.connect() as connection:
            for tabla in metadata.sorted_tables:
                filas = [dict(fila._mapping) for fila in connection.execute(tabla.select())]
                columnas = [
                    columna.name for columna in tabla.columns
                    if self._columna_exportable(columna.name)
                ]
                hojas[tabla.name] = {
                    "columnas": columnas,
                    "filas": [{columna: fila[columna] for columna in columnas} for fila in filas],
                }
        return self._escribir(hojas, destino, datos_crudos=True)

    def _escribir(self, hojas: dict[str, list[dict]], destino: str | Path, datos_crudos: bool = False) -> Path:
        libro = Workbook()
        libro.remove(libro.active)
        for nombre, contenido in hojas.items():
            hoja = libro.create_sheet(nombre[:31])
            if datos_crudos:
                columnas = [columna for columna in contenido["columnas"] if self._columna_exportable(columna)]
                filas = [
                    {columna: fila.get(columna, "") for columna in columnas}
                    for fila in contenido["filas"]
                ]
            else:
                filas = contenido
                columnas = [columna for columna in filas[0] if self._columna_exportable(columna)] if filas else []
            if columnas:
                hoja.append(columnas)
                for celda in hoja[1]:
                    celda.font = Font(bold=True, color="FFFFFF")
                    celda.fill = PatternFill("solid", fgColor="1A1A1A")
                for fila in filas:
                    valores = [fila.get(columna, "") for columna in columnas]
                    celdas = hoja.max_row + 1
                    hoja.append([
                        self._valor_seguro_excel(valor)
                        if isinstance(valor, str) else valor
                        for valor in valores
                    ])
                    for celda in hoja[celdas]:
                        if isinstance(celda.value, str) and celda.value.startswith("'"):
                            celda.data_type = "s"
                hoja.freeze_panes = "A2"
                hoja.auto_filter.ref = hoja.dimensions
                for indice, columna in enumerate(hoja.columns, 1):
                    ancho = min(max(max(len(str(celda.value or "")) for celda in columna) + 2, 12), 48)
                    hoja.column_dimensions[get_column_letter(indice)].width = ancho
        ruta = Path(destino)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        libro.save(ruta)
        return ruta

    @staticmethod
    def _valor_seguro_excel(valor: str) -> str:
        if valor.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
            return f"'{valor}"
        return valor